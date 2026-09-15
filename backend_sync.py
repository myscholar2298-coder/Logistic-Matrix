from datetime import datetime
import io
import os
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path
import openpyxl
import pandas as pd

# ==========================================
# AUTO-DETECT SHARED DRIVE & LOCAL FOLDERS
# ==========================================


def find_shared_drive():
  for drive_letter in ['I', 'H', 'G', 'L', 'D', 'C']:
    test_path = Path(f'{drive_letter}:\\Shared drives\\OPERATION')
    if test_path.exists():
      return test_path
    alt_path = Path(f'{drive_letter}:\\OPERATION')
    if alt_path.exists():
      return alt_path
  return Path(r'H:\Shared drives\OPERATION')  # Default fallback


SHARED_DRIVE_ROOT = find_shared_drive()
LP_FOLDER = SHARED_DRIVE_ROOT / 'LP'

# The git repo is ALWAYS the folder this script lives in.
# (Old version searched for the ancient Google Drive path and could push
#  to the wrong copy of the repo.)
GITHUB_REPO_DIR = Path(__file__).resolve().parent

DELIVERY_PLAN_FOLDER = SHARED_DRIVE_ROOT / 'Delivery Plan'
CUSTOMER_STATEMENT_FOLDER = SHARED_DRIVE_ROOT / 'Customer Statement'

GIT_NAME = 'myscholar2298-coder'
GIT_EMAIL = 'myscholar2298@gmail.com'


def run_git(args, check=True):
  """Run a git command; raise RuntimeError with DECODED output on failure."""
  r = subprocess.run(['git'] + args, capture_output=True)
  if check and r.returncode != 0:
    raise RuntimeError(
        f"git {' '.join(args)} failed (exit {r.returncode}):\n"
        f"{r.stderr.decode(errors='replace')}"
    )
  return r


def ensure_git_identity():
  """Self-heal if git identity is missing."""
  name = run_git(['config', '--get', 'user.name'], check=False)
  email = run_git(['config', '--get', 'user.email'], check=False)
  if not name.stdout.strip() or not email.stdout.strip():
    run_git(['config', 'user.name', GIT_NAME])
    run_git(['config', 'user.email', GIT_EMAIL])
    print(f'Git identity auto-set: {GIT_NAME} / {GIT_EMAIL}')


def normalize_subject(subj_str):
  if not subj_str:
    return ''
  s = str(subj_str).upper().strip()
  return re.sub(r'^PAM', 'PA', s)


def save_csv(df, filename, retries=3, delay_s=10):
  """Write a CSV with retry; survives Excel file locks instead of crashing.
  Returns True on success, False if the file stayed locked."""
  for attempt in range(1, retries + 1):
    try:
      df.to_csv(filename, index=False)
      print(f'Saved: {filename}')
      return True
    except PermissionError:
      print(
          f'Attempt {attempt}/{retries}: {filename} is locked (open in Excel?).'
          f' Retrying in {delay_s}s...'
      )
      time.sleep(delay_s)
    except Exception as e:
      print(f'Error saving {filename}: {e}')
      return False
  print(f'SKIPPED: {filename} - still locked after {retries} attempts.')
  return False


def safe_float(val):
  if val is None:
    return 0.0
  try:
    s = str(val).strip()
    if s.startswith('#') or s.upper() in ['NAN', 'NONE', '']:
      return 0.0
    return float(s.replace(',', '').replace('RM', ''))
  except Exception:
    return 0.0


def run_backend_sync():
  print('--- Starting Full Local Shared Drive Sync ---')

  # HARD FAIL if the source folder is unreachable
  if not LP_FOLDER.exists():
    print(f'FATAL: LP folder not found at {LP_FOLDER}')
    print('Is Google Drive for Desktop running? Is the drive letter correct?')
    sys.exit(1)

  # 1. Load Master Settings Locally from Delivery Plan folder
  #    - Prefer files named exactly "Master_Settings*"
  #    - Exclude Excel lock files (~$...)
  #    - Deduplicate (Windows glob is case-insensitive, so '*master*' and
  #      '*Master*' would return the same file twice)
  all_candidates = list(DELIVERY_PLAN_FOLDER.glob('*aster*ettings*'))
  if not all_candidates:
    all_candidates = list(DELIVERY_PLAN_FOLDER.glob('*master*')) + list(
        DELIVERY_PLAN_FOLDER.glob('*Master*')
    )
  master_files = []
  for f in all_candidates:
    if '~$' in f.name:
      continue
    if f not in master_files:
      master_files.append(f)

  book_prices = {}

  if master_files:
    master_path = master_files[0]
    print(f'Reading Master Settings from: {master_path}')
    try:
      try:
        xls_ms = pd.ExcelFile(master_path, engine='openpyxl')
        sheet_ms = next(
            (
                s
                for s in xls_ms.sheet_names
                if any(
                    k in s.lower()
                    for k in ['invent', 'price', 'harga', 'buku']
                )
            ),
            xls_ms.sheet_names[0],
        )
        print(f'Master sheets found: {xls_ms.sheet_names} -> using: {sheet_ms}')
        df_ms = pd.read_excel(xls_ms, sheet_name=sheet_ms)
      except Exception:
        df_ms = pd.read_csv(master_path)

      # Locate columns by header name when possible; fall back to positions
      col_names = [str(c).strip().lower() for c in df_ms.columns]

      def find_col(*keys):
        for idx, name in enumerate(col_names):
          if any(k in name for k in keys):
            return idx
        return None

      i_subj = find_col('subject', 'subjek')
      i_item = find_col('item', 'buku')
      i_price = find_col('price', 'harga')
      if i_subj is None or i_item is None or i_price is None:
        i_subj, i_item, i_price = 0, 1, 2

      for _, row in df_ms.iterrows():
        raw_subj = row.iloc[i_subj] if i_subj < len(row) else None
        raw_item = row.iloc[i_item] if i_item < len(row) else None
        if (
            raw_subj is None
            or raw_item is None
            or pd.isna(raw_subj)
            or pd.isna(raw_item)
        ):
          continue
        subj = normalize_subject(str(raw_subj))
        item = str(raw_item).strip().upper()
        if not subj or not item or subj.lower() in ('subject', 'nan'):
          continue

        raw_price = row.iloc[i_price] if i_price < len(row) else None
        try:
          price = float(str(raw_price).replace('RM', '').replace(',', '').strip())
          if pd.isna(price) or price <= 0:
            price = 38.0
        except Exception:
          price = 38.0

        book_prices[f'{subj}_{item}'] = price

      print(
          f'Master Settings loaded: {len(book_prices)} price entries'
          f' (sample: {dict(list(book_prices.items())[:3])})'
      )
    except Exception as e:
      print(f'Error reading master settings: {e}')
  else:
    print(
        'Warning: Master settings file not found in'
        f' {DELIVERY_PLAN_FOLDER}'
    )

  # When a GKT/LK title doesn't name the item type, try these in order
  ITEM_PREFERENCE = ['LEMB', 'SPM', 'MTP', 'NOTA']

  def get_price(b_title):
    """Best-effort price lookup; 38.0 only as last resort."""
    t = str(b_title).strip().upper().replace(' ', '_')
    if t in book_prices:
      return book_prices[t]
    parts = t.split('_', 1)
    if len(parts) == 2 and parts[0] and parts[1]:
      subj = normalize_subject(parts[0])
      item = parts[1]
      key = f'{subj}_{item}'
      if key in book_prices:
        return book_prices[key]
      # trailing phase digits: PA1_MTP_2 -> PA1_MTP
      item_core = re.sub(r'_?\d+$', '', item)
      if item_core and item_core != item:
        key = f'{subj}_{item_core}'
        if key in book_prices:
          return book_prices[key]
      # item prefix: SV1_RUJUK -> SV1_RUJ
      for k, v in book_prices.items():
        k_subj, _, k_item = k.partition('_')
        if normalize_subject(k_subj) == subj and item.startswith(k_item):
          return v
    # full-name containment (IPG books): 'SINTAKSIS BAHASA MELAYU' -> SINTAKSIS_RUJ
    title_words = t.replace('_', ' ')
    best_len, best_price = 0, None
    for k, v in book_prices.items():
      ks, _, ki = k.partition('_')
      kt = k.replace('_', ' ')
      score = 0
      if len(kt) >= 4 and (kt in title_words or title_words in kt):
        score = len(kt)
      elif len(ks) >= 5 and ks in title_words:
        score = len(ks)
      elif len(ki) >= 5 and ki in title_words:
        score = len(ki)
      if score > best_len:
        best_len, best_price = score, v
    if best_price is not None:
      return best_price
    # GKT/LK patterns: "GKT T4'26" -> GKT4, "MTP GKT SPM'25" -> GKT+SPM
    words = title_words.replace("'", '')
    m = re.search(r'\b(GKT|LK)\s*T\s?(\d)', words)
    if m:
      base = f'{m.group(1)}{m.group(2)}'
      for kw in ['SPM', 'MTP', 'LEMB', 'NOTA']:
        if kw in words and f'{base}_{kw}' in book_prices:
          return book_prices[f'{base}_{kw}']
      for kw in ITEM_PREFERENCE:
        if f'{base}_{kw}' in book_prices:
          return book_prices[f'{base}_{kw}']
    else:
      fam = next(
          (f for f in ['GKT', 'LK'] if re.search(rf'\b{f}\b', words)), None
      )
      if fam:
        for kw in ['SPM', 'MTP', 'LEMB', 'NOTA']:
          if kw in words:
            hits = [
                v for k, v in book_prices.items()
                if k.startswith(fam) and k.endswith(f'_{kw}')
            ]
            if hits:
              return hits[-1]
    return 38.0

  print('Price self-check:', {
      'PA1 NOTA': get_price('PA1 NOTA'),
      'PA1 MTP 2': get_price('PA1 MTP 2'),
      'SV1 RUJUK': get_price('SV1 RUJUK'),
  })

  # 2. Process Local Sales Status Files in LP Folder (Flexible Matching)
  sales_file_paths = [
      f
      for f in LP_FOLDER.glob('*.xls*')
      if 'sales status' in f.name.lower() and '~$' not in f.name
  ]

  all_sales_records = []
  ledger_records = []
  summary_audit_records = []

  for file_path in sales_file_paths:
    label = (
        file_path.stem.replace(' Sales Status', '')
        .replace('SALES STATUS', '')
        .replace('_', ' ')
        .strip()
    )
    print(f'Processing Sales Status file: {file_path.name} (Label: {label})')

    try:
      with open(file_path, 'rb') as f:
        file_bytes = f.read()

      excel_data = io.BytesIO(file_bytes)
      wb_ox = openpyxl.load_workbook(excel_data, data_only=True)
      excel_data.seek(0)
      xls = pd.ExcelFile(excel_data, engine='openpyxl')

      # Parse Sales Sheets (starting with '^')
      for sheet_name in xls.sheet_names:
        if not str(sheet_name).startswith('^'):
          continue
        df_raw = pd.read_excel(xls, sheet_name=sheet_name, header=None)
        if df_raw.shape[0] < 10:
          continue

        header_row_idx = 9
        for r_idx in range(min(12, len(df_raw))):
          row_text = ' '.join([
              str(x).upper() for x in df_raw.iloc[r_idx].values if pd.notna(x)
          ])
          if 'SCHOOL' in row_text or 'FACULTY' in row_text or 'ROUTE' in row_text:
            header_row_idx = r_idx
            break

        if df_raw.shape[1] > 1:
          df_raw.iloc[:, 1] = df_raw.iloc[:, 1].ffill().astype(str).str.strip()
        if df_raw.shape[1] > 2:
          df_raw.iloc[:, 2] = df_raw.iloc[:, 2].ffill().astype(str).str.strip()

        ws_ox = wb_ox[sheet_name]
        phase_col_info = []
        for r_num in range(8, 13):
          if r_num > ws_ox.max_row:
            continue
          for c_num in range(1, ws_ox.max_column + 1):
            cell_val = ws_ox.cell(row=r_num, column=c_num).value
            if cell_val and str(cell_val).strip().upper() == '1ST':
              title_col_num = c_num - 2
              b_title = ''
              if title_col_num >= 1:
                t_val = ws_ox.cell(row=r_num, column=title_col_num).value
                b_title = str(t_val or '').strip().upper()
              if (
                  b_title
                  and b_title
                  not in ['NAN', 'NONE', '', 'DELIVERED', 'MARKET', 'QUANTITY']
                  and not b_title.startswith('SAMPLE')
              ):
                for p_offset, phase_name in enumerate(
                    ['1ST', '2ND', '3RD', '4TH', '5TH']
                ):
                  target_col_openpyxl = c_num + p_offset
                  if target_col_openpyxl <= ws_ox.max_column:
                    phase_col_info.append(
                        (target_col_openpyxl - 1, phase_name, b_title)
                    )

        for row_idx in range(header_row_idx + 2, len(df_raw)):
          school = df_raw.iloc[row_idx, 1]
          if (
              pd.isna(school)
              or str(school).strip() == ''
              or str(school).strip().upper()
              in ['TOTAL', 'SUBTOTAL', 'BIL', 'NAN']
          ):
            continue
          school_str = str(school).strip()
          teacher = df_raw.iloc[row_idx, 2]
          explicit_debtor = (
              str(df_raw.iloc[row_idx, 4]).strip().upper()
              if df_raw.shape[1] > 4 and pd.notna(df_raw.iloc[row_idx, 4])
              else ''
          )
          if explicit_debtor and not explicit_debtor.startswith('300-'):
            explicit_debtor = f'300-{explicit_debtor}'

          for col_i, sub_h, book_t in phase_col_info:
            if col_i < df_raw.shape[1]:
              val = df_raw.iloc[row_idx, col_i]
              if pd.notna(val) and val != '':
                try:
                  num_val = int(float(val))
                  if num_val != 0:
                    u_price = get_price(book_t)
                    norm_subj = normalize_subject(sheet_name.replace('^', ''))
                    all_sales_records.append({
                        'File_Source': label.strip(),
                        'Subject': norm_subj,
                        'Book_Title': book_t,
                        'School_Name': school_str,
                        'Teacher': (
                            str(teacher).strip()
                            if pd.notna(teacher)
                            and str(teacher).strip().upper() != 'NAN'
                            else 'N/A'
                        ),
                        'Phase': sub_h,
                        'Quantity': num_val,
                        'Unit_Price': u_price,
                        'Total_Value': num_val * u_price,
                        'Explicit_Debtor': explicit_debtor,
                    })
                except Exception:
                  pass

      # Parse Stock Ledger
      excel_data_ledger = io.BytesIO(file_bytes)
      xls_led = pd.ExcelFile(excel_data_ledger, engine='openpyxl')
      target_sheet_led = next(
          (s for s in xls_led.sheet_names if 'stock in' in s.lower()),
          xls_led.sheet_names[0],
      )
      df_raw_led = pd.read_excel(xls_led, sheet_name=target_sheet_led, header=None)
      current_book = ''
      for r_idx in range(len(df_raw_led)):
        row_vals = df_raw_led.iloc[r_idx].values
        col_a = str(row_vals[0]).strip() if pd.notna(row_vals[0]) else ''
        col_b = (
            str(row_vals[1]).strip().upper()
            if len(row_vals) > 1 and pd.notna(row_vals[1])
            else ''
        )
        if col_a and col_a not in ['nan', 'None', '', 'TITLE', 'BIL']:
          current_book = col_a.upper()
        if col_b in ['TOTAL', 'JOHOR', 'MELAKA'] and current_book:
          ledger_records.append({
              'File_Source': label.strip(),
              'Book_Title': current_book,
              'Region': 'TOTAL' if col_b == 'TOTAL' else col_b.capitalize(),
              'Unit_Price': get_price(current_book),
              'Return_%': (
                  safe_float(row_vals[2]) * 100 if len(row_vals) > 2 else 0.0
              ),
              'Total_In': safe_float(row_vals[4]) if len(row_vals) > 4 else 0.0,
              'Total_Return': (
                  safe_float(row_vals[5]) if len(row_vals) > 5 else 0.0
              ),
              'Net_Purchase': (
                  safe_float(row_vals[3]) if len(row_vals) > 3 else 0.0
              ),
          })

      # Parse Stock Summary (With Percentage Fix)
      excel_data_audit = io.BytesIO(file_bytes)
      wb_audit = openpyxl.load_workbook(excel_data_audit, data_only=True)
      target_sheet_name = next(
          (s for s in wb_audit.sheetnames if 'stock summary' in s.lower()), None
      )
      if target_sheet_name:
        ws = wb_audit[target_sheet_name]
        for r in range(3, ws.max_row + 1):
          title_val = ws.cell(row=r, column=1).value
          if not title_val or str(title_val).strip().upper() in [
              'NAN',
              'NONE',
              '',
              'TITLE',
              'BIL',
          ]:
            continue
          book_title = str(title_val).strip().upper()

          job_comp_raw = ws.cell(row=r, column=2).value
          if job_comp_raw is not None:
            try:
              f_val = float(job_comp_raw)
              if f_val >= 0.99 or f_val == 0.01:
                comp_str = '100%'
              elif f_val == 1.0 or f_val == 1:
                comp_str = '100%'
              elif f_val < 1.0:
                comp_str = f'{int(round(f_val * 100))}%'
              else:
                comp_str = f'{int(round(f_val))}%'
            except Exception:
              comp_str = str(job_comp_raw).strip()
          else:
            comp_str = '0%'

          def parse_int(val):
            if val is None:
              return 0
            try:
              return int(float(str(val).replace(',', '').strip()))
            except Exception:
              return 0

          summary_audit_records.append({
              'File_Source': label.strip(),
              'BookType': book_title,
              'Job Completion': comp_str,
              'Sample Balance': parse_int(ws.cell(row=r, column=3).value),
              'Stock Balance': parse_int(ws.cell(row=r, column=4).value),
              'Sample Discrepancy': parse_int(ws.cell(row=r, column=5).value),
              'Stock Discrepancy': parse_int(ws.cell(row=r, column=6).value),
              'Sample On Hand': parse_int(ws.cell(row=r, column=7).value),
              'Stock On Hand': parse_int(ws.cell(row=r, column=8).value),
          })

    except Exception as e:
      print(f'Error processing file {file_path.name}: {e}')

  # 3. Parse Payments and Adjustments Locally from Customer Statement Folder
  pay_path = CUSTOMER_STATEMENT_FOLDER / 'ReceivePayment.xlsx'
  if pay_path.exists():
    try:
      df_pay = pd.read_excel(pay_path, sheet_name=0, engine='openpyxl')
      save_csv(df_pay, 'payments.csv')
      print('Successfully processed local ReceivePayment.xlsx')
    except Exception as e:
      print(f'Error reading local ReceivePayment.xlsx: {e}')

  adj_path = CUSTOMER_STATEMENT_FOLDER / 'Adjustments.xlsx'
  if adj_path.exists():
    try:
      xls_adj = pd.ExcelFile(adj_path, engine='openpyxl')
      adj_sheet = next(
          (
              s
              for s in xls_adj.sheet_names
              if 'credit note' in s.lower() or 'adj' in s.lower()
          ),
          xls_adj.sheet_names[0],
      )
      df_adj = pd.read_excel(xls_adj, sheet_name=adj_sheet)
      save_csv(df_adj, 'adjustments.csv')
      print('Successfully processed local Adjustments.xlsx')
    except Exception as e:
      print(f'Error reading local Adjustments.xlsx: {e}')

  # Save primary operational CSVs
  df_sales_final = pd.DataFrame(all_sales_records)
  save_ok = True
  save_ok &= save_csv(df_sales_final, 'sales_transactions.csv')
  save_ok &= save_csv(pd.DataFrame(ledger_records), 'stock_ledger.csv')
  save_ok &= save_csv(pd.DataFrame(summary_audit_records), 'stock_summary.csv')

  # 4. Pre-calculate Menu 2 Outstanding Summary
  print('Pre-calculating Menu 2 Outstanding Summary...')
  if not df_sales_final.empty:
    summary_records = []
    for (file_src, teacher_name, debtor_code), group_df in df_sales_final.groupby(
        ['File_Source', 'Teacher', 'Explicit_Debtor']
    ):
      group_deliveries_val = group_df['Total_Value'].sum()
      group_qty = group_df['Quantity'].sum()
      group_count = len(group_df)
      schools_in_group = group_df['School_Name'].unique().tolist()
      primary_school = schools_in_group[0] if schools_in_group else ''
      primary_subject = normalize_subject(group_df['Subject'].iloc[0])

      summary_records.append({
          'File Source': file_src,
          'Subject': primary_subject,
          'Debtor Code': debtor_code if pd.notna(debtor_code) else 'N/A',
          'Teacher': teacher_name if teacher_name and teacher_name != 'N/A' else 'N/A',
          'Associated School(s)': ', '.join(schools_in_group),
          'Net Qty': group_qty,
          'Outstanding (RM)': group_deliveries_val,
          'Transactions': group_count,
          '_school': primary_school,
          '_subject': primary_subject,
          '_teacher': teacher_name,
      })

    save_ok &= save_csv(pd.DataFrame(summary_records), 'outstanding_summary.csv')

  if not save_ok:
    print('FATAL: One or more CSVs could not be saved (still locked?).')
    print('The data pushed to GitHub would be stale - aborting push.')
    sys.exit(1)

  # 5. Push BI CSVs + Master_Settings to GitHub (same repo as dispatch CSV)
  push_files = [
      'sales_transactions.csv',
      'stock_ledger.csv',
      'stock_summary.csv',
      'outstanding_summary.csv',
      'payments.csv',
      'adjustments.csv',
  ]
  try:
    copied = []
    for fname in push_files:
      src = Path.cwd() / fname
      if src.exists():
        dst = GITHUB_REPO_DIR / fname
        if src.resolve() != dst.resolve():
          shutil.copy(src, dst)
        copied.append(fname)
    if master_files:
      shutil.copy(master_files[0], GITHUB_REPO_DIR / 'Master_Settings.xlsx')
      copied.append('Master_Settings.xlsx')
    if not copied:
      print('GitHub push: no files to push.')
    else:
      os.chdir(GITHUB_REPO_DIR)
      ensure_git_identity()
      run_git(['add'] + copied)
      made_commit = False
      if run_git(['diff', '--cached', '--quiet'], check=False).returncode != 0:
        run_git([
            'commit', '-m',
            f'BI data sync {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}',
        ])
        made_commit = True
      # ALWAYS push - even if this script committed nothing, AutoUpdate.py
      # may have left an unpushed commit that only this push can deliver.
      try:
        run_git(
            ['pull', '--rebase', '--autostash', '-X', 'theirs', 'origin', 'main']
        )
      except RuntimeError:
        run_git(['rebase', '--abort'], check=False)
        raise
      run_git(['push', 'origin', 'main'])
      if made_commit:
        print(f'GitHub push: {len(copied)} files pushed.')
      else:
        print('GitHub push: no new BI commit, but pending commits were pushed.')
  except Exception as e:
    # LOUD failure - Task Scheduler will show a non-zero result
    print(f'FATAL: GitHub push failed:\n{e}')
    sys.exit(1)

  print('--- Local Shared Drive Sync Complete! ---')


if __name__ == '__main__':
  run_backend_sync()